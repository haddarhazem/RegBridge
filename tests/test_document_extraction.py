from types import SimpleNamespace

import pytest
from docx import Document as DocxDocument
from fpdf import FPDF

from app.modules.documents.contract_analysis_service import ContractAnalysisService
from app.modules.documents.extraction import ExtractionError, extract_native
from app.modules.documents.ocr import MistralOcrProvider


def test_utf8_text_is_extracted_without_ocr(tmp_path):
    source = tmp_path / "notice.txt"
    source.write_bytes("Première ligne\r\nDeuxième ligne".encode("utf-8"))

    result = extract_native(source, ".txt")

    assert result.method == "native"
    assert result.extractor_version == "native-text-v1"
    assert result.normalized_text == "Première ligne\nDeuxième ligne"
    assert result.pages == (result.normalized_text,)


def test_docx_extraction_preserves_body_and_table_order(tmp_path):
    source = tmp_path / "terms.docx"
    document = DocxDocument()
    document.add_paragraph("First clause")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Party"
    table.cell(0, 1).text = "Obligation"
    document.add_paragraph("Last clause")
    document.save(source)

    result = extract_native(source, ".docx")

    assert result.normalized_text == "First clause\nParty\tObligation\nLast clause"


def test_machine_readable_pdf_keeps_page_boundary(tmp_path):
    source = tmp_path / "notice.pdf"
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.multi_cell(0, 10, "A readable regulatory obligation")
    pdf.output(str(source))

    result = extract_native(source, ".pdf", min_pdf_chars=8)

    assert result.method == "native"
    assert result.pages and "readable regulatory obligation" in result.pages[0]
    assert result.normalized_text.startswith("PAGE 1\n")


def test_image_only_pdf_requires_opt_in_ocr(tmp_path):
    source = tmp_path / "scan.pdf"
    pdf = FPDF()
    pdf.add_page()
    pdf.output(str(source))

    with pytest.raises(ExtractionError, match="no sufficient native text") as error:
        extract_native(source, ".pdf", min_pdf_chars=8)

    assert error.value.category == "OCR_REQUIRED"


def test_mixed_pdf_retains_native_page_and_marks_only_unusable_page_for_ocr(tmp_path):
    source = tmp_path / "mixed.pdf"
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.multi_cell(0, 10, "This native page has enough selectable contract text to remain locally extracted.")
    pdf.add_page()
    pdf.output(str(source))

    result = extract_native(source, ".pdf", min_pdf_chars=16)

    assert result.method == "native"
    assert result.pages[0].startswith("This native page")
    assert result.ocr_required_pages == (2,)
    assert result.page_methods == ("NATIVE", "OCR_REQUIRED")


@pytest.mark.asyncio
async def test_mistral_ocr_uploads_exact_bytes_and_deletes_provider_file():
    calls = []

    class Files:
        async def upload_async(self, **kwargs):
            calls.append(("upload", kwargs))
            return SimpleNamespace(id="provider-file-id")

        async def delete_async(self, **kwargs):
            calls.append(("delete", kwargs))

    class OCR:
        async def process_async(self, **kwargs):
            calls.append(("ocr", kwargs))
            return SimpleNamespace(model="mistral-ocr-test", pages=[SimpleNamespace(markdown="Scanned obligation")])

    provider = MistralOcrProvider(api_key="test-only-key", model="mistral-ocr-test", client=SimpleNamespace(files=Files(), ocr=OCR()))

    result = await provider.extract(b"private-pdf-bytes", "scan.pdf", "application/pdf")

    assert result.method == "mistral_ocr"
    assert result.provider == "mistral"
    assert result.model == "mistral-ocr-test"
    assert result.normalized_text == "PAGE 1\nScanned obligation"
    assert [call[0] for call in calls] == ["upload", "ocr", "delete"]
    assert calls[0][1]["purpose"] == "ocr"
    assert calls[0][1]["visibility"] == "user"
    assert calls[0][1]["file"].content == b"private-pdf-bytes"
    assert calls[1][1]["document"].file_id == "provider-file-id"
    assert calls[2][1] == {"file_id": "provider-file-id", "timeout_ms": None}


@pytest.mark.asyncio
async def test_mistral_ocr_invalid_response_is_categorized_and_provider_file_is_removed():
    calls = []

    class Files:
        async def upload_async(self, **_kwargs):
            return SimpleNamespace(id="provider-file-id")

        async def delete_async(self, **kwargs):
            calls.append(kwargs)

    class OCR:
        async def process_async(self, **_kwargs):
            return SimpleNamespace(model="mistral-ocr-test", pages=[])

    provider = MistralOcrProvider(
        api_key="test-only-key",
        model="mistral-ocr-test",
        client=SimpleNamespace(files=Files(), ocr=OCR()),
    )

    with pytest.raises(ExtractionError) as error:
        await provider.extract(b"private-pdf-bytes", "scan.pdf", "application/pdf")

    assert error.value.category == "OCR_INVALID_RESPONSE"
    assert calls == [{"file_id": "provider-file-id", "timeout_ms": None}]


def test_contract_analysis_source_method_uses_page_level_mixed_provenance():
    version = SimpleNamespace(
        extraction_metadata={
            "extraction": {
                "method": "hybrid_native_ocr",
                "page_methods": ["NATIVE", "OCR"],
            }
        }
    )

    assert ContractAnalysisService._source_method(version) == "MIXED"
